"""字体规范：标题黑体 / 正文宋体 / 西文 Times New Roman（project note-论文撰写-字体）。

两半：
- `deck_fonts.py` 按字号分角色写 `a:latin` / `a:ea`（≥30pt 算标题）；
- `deck_check.py` 查字体族是否在规范表、中文有没有写 `a:ea`、角色是否一致。
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import deck_check  # noqa: E402
import deck_fonts  # noqa: E402

W = 13.333


def new_deck():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(7.5)
    return prs


def add_text(slide, text, size, x=0.8, y=0.8, w=9.0, h=1.2, font=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.text = text
    for run in box.text_frame.paragraphs[0].runs:
        run.font.size = Pt(size)
        if font:
            run.font.name = font
    return box


def add_note(slide):
    slide.notes_slide.notes_text_frame.text = "本页口播：先结论后证据。"


def test_fonts_role_split(tmp_path):
    prs = new_deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(slide, "可用率做到 96%", 54)
    add_text(slide, "正文说明一句话。", 18, y=2.5)
    add_note(slide)
    path = tmp_path / "fonts.pptx"
    prs.save(path)

    result = deck_fonts.deal_fonts(path, apply=True)
    assert result["title"] == 1 and result["body"] == 1
    assert result["changed"] == 2

    prs2 = Presentation(str(path))
    runs = [run for shape in prs2.slides[0].shapes
            if shape.has_text_frame for para in shape.text_frame.paragraphs
            for run in para.runs]
    latin, title_ea = deck_check._run_typefaces(runs[0])
    _, body_ea = deck_check._run_typefaces(runs[1])
    assert latin == "Times New Roman"
    assert title_ea == "SimHei"
    assert body_ea == "SimSun"


def test_missing_ea_is_warning(tmp_path):
    prs = new_deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(slide, "这行中文没有写 a:ea", 18, font="Times New Roman")
    add_note(slide)
    path = tmp_path / "noea.pptx"
    prs.save(path)

    report = deck_check.check_deck(path)
    assert any(i.code == "font-ea-missing" for i in report.warnings)


def test_role_mismatch_is_warning(tmp_path):
    prs = new_deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(slide, "标题用了宋体", 54)
    add_note(slide)
    path = tmp_path / "role.pptx"
    prs.save(path)
    deck_fonts.deal_fonts(path, title_ea="SimSun", body_ea="SimSun", apply=True)

    report = deck_check.check_deck(path, title_ea="SimHei", body_ea="SimSun")
    assert any(i.code == "font-role" for i in report.warnings)


def test_unknown_face_is_warning(tmp_path):
    prs = new_deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(slide, "字体不合规", 18, font="Comic Sans MS")
    add_note(slide)
    path = tmp_path / "face.pptx"
    prs.save(path)

    report = deck_check.check_deck(path)
    assert any(i.code == "font-face" for i in report.warnings)
