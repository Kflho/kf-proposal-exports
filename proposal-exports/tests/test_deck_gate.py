"""deck_gate.py：一条命令跑完页片闸门（本仓检查 + 外部审计接线上）。

判据：
- 全部步骤 exit 0 才算通过；外部脚本找不到时标 skipped，不算失败；
- 参数错误（未知步骤、显式指定的外部脚本不存在）返回 2；
- 任一被调脚本非 0，闸门返回 1。
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import deck_gate  # noqa: E402

W = 13.333


def make_deck(path: Path, ratio: float = 16 / 9) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(W / ratio)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(9), Inches(1.2))
    box.text_frame.text = "多站协同定位把测距误差压到 10 m 以内"
    for run in box.text_frame.paragraphs[0].runs:
        run.font.size = Pt(32)
    slide.notes_slide.notes_text_frame.text = "先给结论，再讲证据与限定条件。"
    prs.save(path)
    return path


def test_find_external_audit_explicit(tmp_path):
    script = tmp_path / "audit.py"
    script.write_text("raise SystemExit(0)\n", encoding="utf-8")
    assert deck_gate.find_external_audit(str(script)) == script
    assert deck_gate.find_external_audit(str(tmp_path / "missing.py")) is None


def test_find_external_audit_env(monkeypatch, tmp_path):
    script = tmp_path / "audit.py"
    script.write_text("raise SystemExit(0)\n", encoding="utf-8")
    monkeypatch.setenv("KF_DECK_AUDIT", str(script))
    assert deck_gate.find_external_audit(None, candidates=()) == script


def test_gate_passes_clean_deck(tmp_path):
    deck = make_deck(tmp_path / "ok.pptx")
    assert deck_gate.main([str(deck), "--steps", "deck_check"]) == 0


def test_gate_fails_on_bad_aspect(tmp_path):
    deck = make_deck(tmp_path / "bad.pptx", ratio=4 / 3)
    assert deck_gate.main([str(deck), "--steps", "deck_check"]) == 1


def test_gate_skips_missing_external(tmp_path, monkeypatch):
    deck = make_deck(tmp_path / "ok.pptx")
    monkeypatch.setattr(deck_gate, "find_external_audit", lambda *_args, **_kwargs: None)
    assert deck_gate.main([str(deck), "--steps", "external"]) == 0


def test_gate_fails_when_external_fails(tmp_path):
    deck = make_deck(tmp_path / "ok.pptx")
    script = tmp_path / "audit.py"
    script.write_text("raise SystemExit(1)\n", encoding="utf-8")
    code = deck_gate.main([str(deck), "--steps", "external", "--external-script", str(script)])
    assert code == 1


def test_gate_rejects_unknown_step(tmp_path):
    deck = make_deck(tmp_path / "ok.pptx")
    assert deck_gate.main([str(deck), "--steps", "nope"]) == 2
