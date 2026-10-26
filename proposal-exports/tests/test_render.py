"""render_pptx_png：预览要么出 PNG，要么明确报"需要 Windows + PowerPoint"。

不能静默降级：没有 PowerPoint 时返回非 0 并在消息里说清缺什么。
"""
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "render_pptx_png.py"


def make_deck(path: Path) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = "预览冒烟"
    prs.save(path)
    return path


def test_render_or_explicit_error(tmp_path):
    deck = make_deck(tmp_path / "deck.pptx")
    out = tmp_path / "preview"
    proc = subprocess.run([sys.executable, str(SCRIPT), str(deck), "--out", str(out)],
                          capture_output=True, text=True, encoding="utf-8",
                          errors="replace", timeout=180)
    if proc.returncode == 0:
        pngs = list(out.glob("*.png"))
        assert pngs, "返回 0 却没有产出 PNG"
    else:
        message = (proc.stdout + proc.stderr).lower()
        assert "powerpoint" in message or "pywin32" in message or "office" in message, \
            f"失败时要说清缺什么：{proc.stdout[-200:]} {proc.stderr[-200:]}"
