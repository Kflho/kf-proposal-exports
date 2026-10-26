"""text_check.py 的行为测试：判据来自 references/text-quality.md。

- 禁用词 = error（"赋能/闭环/里程碑"这类 AI 高频词）；
- "模型""结果分析"这类话题式标题 = warning（标题要写结论）；
- 长句/要点过多/口播过长 = warning。
"""
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import text_check  # noqa: E402


def write_md(tmp_path, text):
    path = tmp_path / "doc.md"
    path.write_text(text, encoding="utf-8")
    return path


def test_banned_phrase_is_error(tmp_path):
    path = write_md(tmp_path, "## 方案\n\n本方案可以赋能整条产业链。\n")
    report = text_check.check_file(path)
    assert any(i.code == "banned-phrase" for i in report.errors)
    assert not report.ok


def test_topic_style_title_is_warning(tmp_path):
    path = write_md(tmp_path, "## 模型\n\n这里介绍方程。\n")
    report = text_check.check_file(path)
    assert any(i.code == "title-topic" for i in report.warnings)
    assert report.ok


def test_action_title_passes(tmp_path):
    path = write_md(tmp_path, "## 只开径向发动机时秩为 3，切向漂移拉不回来\n\n正文一句话。\n")
    report = text_check.check_file(path)
    assert not [i for i in report.warnings if i.code == "title-topic"]


def test_long_sentence_is_warning(tmp_path):
    long_line = "这是一个非常长的句子" * 8 + "。"
    path = write_md(tmp_path, f"## 一句合格的结论标题\n\n{long_line}\n")
    report = text_check.check_file(path)
    assert any(i.code == "sentence-too-long" for i in report.warnings)


def test_colloquial_register_is_warning(tmp_path):
    path = write_md(tmp_path, "## 单站测距遇遮挡就掉点\n\n站数加到 12 个也补不回。\n")
    report = text_check.check_file(path)
    hits = [i for i in report.warnings if i.code == "colloquial"]
    assert hits, report.warnings
    assert any("掉点" in item.message for item in hits)


def test_academic_register_passes(tmp_path):
    path = write_md(tmp_path, "## 遮挡 40% 时可用率降至 89%\n\n仅增加站数无法恢复可用率。\n")
    report = text_check.check_file(path)
    assert not [i for i in report.warnings if i.code == "colloquial"]


def _deck_with(tmp_path, bullets, notes=""):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(9), Inches(1))
    title.text_frame.text = "同一组极点，两种过渡过程"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(30)
    body = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(9), Inches(3))
    body.text_frame.text = "\n".join(bullets)
    for para in body.text_frame.paragraphs:
        para.runs[0].font.size = Pt(16)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    path = tmp_path / "deck.pptx"
    prs.save(path)
    return path


def test_too_many_bullets_is_warning(tmp_path):
    path = _deck_with(tmp_path, [f"第{i}条要点" for i in range(6)])
    report = text_check.check_file(path)
    assert any(i.code == "too-many-bullets" for i in report.warnings)


def test_long_notes_is_warning(tmp_path):
    path = _deck_with(tmp_path, ["要点一"], notes="口播" * 80)
    report = text_check.check_file(path)
    assert any(i.code == "notes-too-long" for i in report.warnings)


def test_clean_deck_passes(tmp_path):
    path = _deck_with(tmp_path, ["K₁：切向位置荡出 0.63 m"], notes="先讲结论，再指图。")
    report = text_check.check_file(path)
    assert report.ok
    assert report.pages == 1


def test_title_without_verdict_is_warning(tmp_path):
    path = write_md(tmp_path, "## 系统结构与实验安排\n\n正文一句话。\n")
    report = text_check.check_file(path)
    assert any(i.code == "title-no-verdict" for i in report.warnings)


def test_title_with_verdict_passes(tmp_path):
    path = write_md(tmp_path, "## 方案可行，代价是一次迭代\n\n正文一句话。\n")
    report = text_check.check_file(path)
    assert not [i for i in report.warnings if i.code == "title-no-verdict"]


def test_ai_phrase_is_warning(tmp_path):
    path = write_md(tmp_path, "## 缓存把查询压到 5 ms\n\n通过引入缓存实现查询加速。\n")
    report = text_check.check_file(path)
    assert any(i.code == "ai-phrase" for i in report.warnings)


def test_banned_word_in_title_is_error(tmp_path):
    path = write_md(tmp_path, "## 全链路方案\n\n正文一句话。\n")
    report = text_check.check_file(path)
    assert any(i.code == "banned-phrase" for i in report.errors)


def test_allow_context_passes(tmp_path):
    path = write_md(tmp_path, "## 控制闭环保持稳定\n\n控制闭环保持稳定。\n")
    report = text_check.check_file(path)
    assert not [i for i in report.errors if i.code == "banned-phrase"]


def test_control_theory_contexts_pass(tmp_path):
    """闭环传递矩阵/闭环特征多项式是控制术语，不是 AI 高频词。"""
    text = ("## 闭环传递矩阵与题目要求一致\n\n"
            "闭环传递矩阵逐点一致，闭环特征多项式等于 gamma(s)。\n")
    report = text_check.check_file(write_md(tmp_path, text))
    assert not [i for i in report.errors if i.code == "banned-phrase"]


def test_plain_closed_loop_still_errors(tmp_path):
    """去掉固定搭配后，裸「闭环」仍按 AI 高频词报错。"""
    path = write_md(tmp_path, "## 形成闭环\n\n形成闭环。\n")
    report = text_check.check_file(path)
    assert any(i.code == "banned-phrase" for i in report.errors)
