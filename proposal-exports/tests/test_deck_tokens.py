"""页片令牌检查（design-notes/deck-aesthetics.md）：字号档位、强调色预算、不用渐变/阴影。

期望值来自设计记录，不从实现里抄：
- 一页最多五档字号；
- 投影最小 8pt；
- 强调色每页最多两处（结论数字 + 一处框线/引线）；
- 渐变、阴影一律 error（用平色与色阶分层）。
"""
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import deck_check  # noqa: E402

W = 13.333
ACCENT = RGBColor(0xE8, 0xA3, 0x3D)


def new_deck():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(7.5)
    return prs


def add_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.notes_slide.notes_text_frame.text = "这一页讲结论与证据的对应关系。"
    return slide


def add_runs(slide, sizes, x=0.8, y=0.9, color=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(8), Inches(2.4))
    frame = box.text_frame
    frame.text = ""
    para = frame.paragraphs[0]
    for index, size in enumerate(sizes):
        run = para.add_run()
        run.text = f"第{index + 1}条"
        run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
    return box


def inject(slide, shape, tag):
    etree.SubElement(shape._element.spPr, qn(tag))


def codes(report, level=None):
    issues = report.errors + report.warnings
    return [issue.code for issue in issues if level is None or issue.level == level]


def test_gradient_is_error(tmp_path):
    prs = new_deck()
    slide = add_slide(prs)
    box = add_runs(slide, [18])
    inject(slide, box, "a:gradFill")
    path = tmp_path / "gradient.pptx"
    prs.save(path)

    report = deck_check.check_deck(path)
    assert "gradient-fill" in codes(report, "error")


def test_shadow_is_error(tmp_path):
    prs = new_deck()
    slide = add_slide(prs)
    box = add_runs(slide, [18])
    inject(slide, box, "a:effectLst")
    etree.SubElement(box._element.spPr.find(qn("a:effectLst")), qn("a:outerShdw"))
    path = tmp_path / "shadow.pptx"
    prs.save(path)

    report = deck_check.check_deck(path)
    assert "shadow-effect" in codes(report, "error")


def test_accent_over_budget_warns_then_errors(tmp_path):
    prs = new_deck()
    slide = add_slide(prs)
    add_runs(slide, [18, 18, 18], color=ACCENT)
    path = tmp_path / "accent.pptx"
    prs.save(path)

    report = deck_check.check_deck(path)
    assert "accent-budget" in codes(report, "warning")
    strict = deck_check.check_deck(path, strict=True)
    assert "accent-budget" in codes(strict, "error")


def test_type_below_8pt_is_error(tmp_path):
    prs = new_deck()
    slide = add_slide(prs)
    add_runs(slide, [18, 6])
    path = tmp_path / "tiny.pptx"
    prs.save(path)

    report = deck_check.check_deck(path)
    assert "type-too-small" in codes(report, "error")


def test_more_than_five_sizes_warns_and_scale_flags_off_ladder(tmp_path):
    prs = new_deck()
    slide = add_slide(prs)
    add_runs(slide, [54, 33, 24, 20, 18, 12])
    path = tmp_path / "ladder.pptx"
    prs.save(path)

    report = deck_check.check_deck(path, type_scale=[54, 33, 18, 12, 10])
    assert "type-scale" in codes(report, "warning")
    messages = [issue.message for issue in report.warnings if issue.code == "type-scale"]
    assert any("不在档位" in message for message in messages)
