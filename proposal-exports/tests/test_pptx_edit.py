"""update_pptx_images / fit_pptx_frames：换图不毁版式、按真实比例重排图片框。

判据（references/ppt-deck.md）：
- 换图：只替换指定媒体条目，其余条目与形状几何不动；
- 重排框：框的宽高比与图片真实比例一致（容差 3%），且留在安全区内。
"""
import subprocess
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

ROOT = Path(__file__).resolve().parents[1]
UPDATE = ROOT / "scripts" / "update_pptx_images.py"
FIT = ROOT / "scripts" / "fit_pptx_frames.py"


def run(script: Path, *args: str):
    return subprocess.run([sys.executable, str(script), *args], capture_output=True,
                          text=True, encoding="utf-8", errors="replace", timeout=120)

W, H = 13.333, 7.5


def png(path: Path, color, size=(120, 80)) -> Path:
    Image.new("RGB", size, color).save(path)
    return path


def make_deck(path: Path, first, second) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # 故意拉伸：原图 3:2，框 4:2
    slide.shapes.add_picture(str(first), Inches(1), Inches(1), width=Inches(4), height=Inches(2))
    slide.shapes.add_picture(str(second), Inches(6), Inches(1), width=Inches(3), height=Inches(2))
    prs.save(path)
    return path


def media_entries(deck: Path) -> list[str]:
    proc = run(UPDATE, str(deck), "--list")
    assert proc.returncode == 0, proc.stderr
    entries = []
    for line in proc.stdout.splitlines():
        token = line.strip().split()[0] if line.strip() else ""
        if token.startswith("ppt/media/"):
            entries.append(token)
    return entries


def test_update_images_replaces_only_target(tmp_path):
    first = png(tmp_path / "a.png", (200, 30, 30))
    second = png(tmp_path / "b.png", (30, 30, 200))
    deck = make_deck(tmp_path / "deck.pptx", first, second)
    entries = media_entries(deck)
    assert len(entries) >= 2, entries

    new = png(tmp_path / "c.png", (30, 200, 30))
    new_bytes = new.read_bytes()
    proc = run(UPDATE, str(deck), "--map", f"{entries[0]}={new}", "--apply")
    assert proc.returncode == 0, proc.stderr

    blobs = [shape.image.blob for slide in Presentation(str(deck)).slides
             for shape in slide.shapes if shape.shape_type == 13]
    assert new_bytes in blobs
    assert any(blob != new_bytes for blob in blobs), "另一张图不该被替换"


def test_fit_frames_preserves_real_aspect(tmp_path):
    first = png(tmp_path / "a.png", (200, 30, 30))
    second = png(tmp_path / "b.png", (30, 30, 200))
    deck = make_deck(tmp_path / "deck.pptx", first, second)
    proc = run(FIT, str(deck), "--apply")
    assert proc.returncode == 0, proc.stderr

    prs = Presentation(str(deck))
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type != 13:
                continue
            native_w, native_h = shape.image.size
            ratio = (shape.width / shape.height) / (native_w / native_h)
            assert abs(ratio - 1) < 0.03, f"比例仍偏 {ratio:.3f}"
            assert shape.left >= 0 and shape.left + shape.width <= prs.slide_width
            assert shape.top >= 0 and shape.top + shape.height <= prs.slide_height
