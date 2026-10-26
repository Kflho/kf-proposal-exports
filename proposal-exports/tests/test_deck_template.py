"""deck_template.py：抽模板令牌、按版式/复制页生成。

判据：
- tokens 里有画布尺寸、主题色、主题字体、版式占位符；
- plan 按 layout 填标题/正文，按 copy 复制页并替换文字；
- 未命中项要报出来（返回码 1），不能静默丢内容。
"""
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import deck_template  # noqa: E402


def make_template(path: Path) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    sample = prs.slides.add_slide(prs.slide_layouts[1])
    sample.shapes.title.text = "模板样例句"
    body = sample.placeholders[1].text_frame
    body.text = "样例要点一"
    for para in body.paragraphs:
        for run in para.runs:
            run.font.size = Pt(20)
    sample.notes_slide.notes_text_frame.text = "样例页口播，够十个字。"
    prs.save(path)
    return path


def test_extract_tokens(tmp_path):
    template = make_template(tmp_path / "tpl.pptx")
    tokens = deck_template.extract(Presentation(str(template)))
    assert tokens["slide_size"][0] > 13
    assert tokens["theme"]["colors"], "主题色没抽到"
    assert any(layout["placeholders"] for layout in tokens["layouts"]), "版式占位符为空"
    assert tokens["slides"], "页统计为空"


def test_tokens_cli_writes_file(tmp_path):
    template = make_template(tmp_path / "tpl.pptx")
    out = tmp_path / "tokens.json"
    code = deck_template.main([str(template), "--tokens", str(out)])
    assert code == 0
    data = json.loads(out.read_text(encoding="utf-8"))
    assert data["slide_size"] and data["layouts"]


def test_build_from_layout(tmp_path):
    template = make_template(tmp_path / "tpl.pptx")
    plan = {"pages": [
        {"layout": "Title and Content", "title": "遮挡 30% 时可用率 96%",
         "body": ["较基线高 18 个百分点", "求解时间 5 ms"], "notes": "这页的口播够十个字。"},
    ]}
    out = tmp_path / "out.pptx"
    result = deck_template.build(template, plan, out)
    assert result["unresolved"] == []
    prs = Presentation(str(out))
    assert len(prs.slides) == 1  # 模板样例页被清掉
    texts = [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]
    assert any("可用率 96%" in t for t in texts)
    assert any("18 个百分点" in t for t in texts)
    assert prs.slides[0].notes_slide.notes_text_frame.text.startswith("这页")


def test_build_copy_page_replaces_text(tmp_path):
    template = make_template(tmp_path / "tpl.pptx")
    plan = {"pages": [
        {"copy": 1, "text": {"模板样例句": "复制页新标题"}, "notes": "复制页口播也够十个字。"},
    ]}
    out = tmp_path / "copy.pptx"
    result = deck_template.build(template, plan, out)
    assert result["unresolved"] == []
    prs = Presentation(str(out))
    texts = [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]
    assert any(t == "复制页新标题" for t in texts)


def test_unresolved_is_reported(tmp_path):
    template = make_template(tmp_path / "tpl.pptx")
    plan = {"pages": [{"copy": 1, "text": {"不存在的原句": "新句"}}]}
    result = deck_template.build(template, plan, tmp_path / "bad.pptx")
    assert result["unresolved"], "未命中的替换目标必须报出来"
