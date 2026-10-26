"""deck_spec.py：把 pptx 拆成 spec、比两版差异（页片回灌的依据）。

判据：
- dump 保留页数、文字、表格、备注；
- 两版完全一致 → 0 差异；
- 文字变了 → 报 added + removed 一对；
- 备注变了 → 报 notes。
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import deck_spec  # noqa: E402


def make_deck(path: Path, text: str = "可用率做到 96%", notes: str = "开场给结论。") -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(8), Inches(1))
    box.text_frame.text = text
    for run in box.text_frame.paragraphs[0].runs:
        run.font.size = Pt(30)
    slide.notes_slide.notes_text_frame.text = notes

    second = prs.slides.add_slide(prs.slide_layouts[6])
    table = second.shapes.add_table(2, 2, Inches(0.8), Inches(1.2), Inches(6), Inches(1.2)).table
    table.cell(0, 0).text = "工况"
    table.cell(0, 1).text = "秩"
    table.cell(1, 0).text = "仅径向"
    table.cell(1, 1).text = "3"
    second.notes_slide.notes_text_frame.text = "第二页讲能控性。"
    prs.save(path)
    return path


def test_dump_keeps_structure(tmp_path):
    deck = make_deck(tmp_path / "a.pptx")
    spec = deck_spec.dump_deck(deck)
    assert spec["pages"] == 2
    kinds = [item["kind"] for item in spec["slides"][0]["items"]]
    assert "text" in kinds
    kinds2 = [item["kind"] for item in spec["slides"][1]["items"]]
    assert "table" in kinds2
    assert spec["slides"][1]["notes"].startswith("第二页")


def test_compare_identical_is_clean(tmp_path):
    a = make_deck(tmp_path / "a.pptx")
    result = deck_spec.compare_specs(deck_spec.dump_deck(a), deck_spec.dump_deck(a))
    assert result["count"] == 0


def test_compare_text_change(tmp_path):
    a = make_deck(tmp_path / "a.pptx")
    b = make_deck(tmp_path / "b.pptx", text="可用率做到 97%")
    result = deck_spec.compare_specs(deck_spec.dump_deck(a), deck_spec.dump_deck(b))
    kinds = {item["kind"] for item in result["diffs"]}
    assert {"added", "removed"} <= kinds


def test_compare_notes_change(tmp_path):
    a = make_deck(tmp_path / "a.pptx")
    b = make_deck(tmp_path / "b.pptx", notes="换一句口播。")
    result = deck_spec.compare_specs(deck_spec.dump_deck(a), deck_spec.dump_deck(b))
    assert any(item["kind"] == "notes" for item in result["diffs"])
