"""deck_critique.py：审美评估的机械法条（L1–L8）。

判据来自 references/deck-critique.md：
- L1 一页 ≤5 档字号、最小 ≥8pt（FAIL）
- L2 强调色每页 ≤2 处（FAIL）
- L3 无渐变/阴影（FAIL）
- L5 标题与正文左缘对齐（REVIEW）
"""
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import deck_critique  # noqa: E402

ACCENT = RGBColor(0xE8, 0xA3, 0x3D)


def new_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def add_text(slide, text, size, x=0.8, y=0.8, w=9.0, h=1.0, center=False, color=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.text = text
    para = box.text_frame.paragraphs[0]
    if center:
        from pptx.enum.text import PP_ALIGN
        para.alignment = PP_ALIGN.CENTER
    for run in para.runs:
        run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
    return box


def add_filler(slide):
    """放一个表格撑起覆盖率，避免每页都判"偏空"。"""
    table = slide.shapes.add_table(3, 3, Inches(0.8), Inches(2.4), Inches(8), Inches(3.6)).table
    for row in range(3):
        for col in range(3):
            table.cell(row, col).text = "数据"


def codes(report):
    return [item.rule for item in report.findings if item.level == "FAIL"]


def test_clean_deck_has_no_fail(tmp_path):
    prs = new_deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(slide, "遮挡 30% 时可用率 96%", 30)
    add_text(slide, "较基线高 18 个百分点", 15, y=1.9)
    add_filler(slide)
    path = tmp_path / "ok.pptx"
    prs.save(path)

    report = deck_critique.critique(path)
    assert not report.failed, [item for item in report.findings if item.level == "FAIL"]


def test_small_font_fails(tmp_path):
    prs = new_deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(slide, "正文", 6)
    add_filler(slide)
    path = tmp_path / "small.pptx"
    prs.save(path)
    assert "L1 字号档位" in codes(deck_critique.critique(path))


def test_gradient_fails(tmp_path):
    prs = new_deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = add_text(slide, "标题", 30)
    add_filler(slide)
    etree.SubElement(box._element.spPr, qn("a:gradFill"))
    path = tmp_path / "grad.pptx"
    prs.save(path)
    assert "L3 平色" in codes(deck_critique.critique(path))


def test_accent_over_budget_fails(tmp_path):
    prs = new_deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(slide, "结论一", 15, color=ACCENT)
    add_text(slide, "结论二", 15, y=2.0, color=ACCENT)
    add_text(slide, "结论三", 15, y=3.0, color=ACCENT)
    add_filler(slide)
    path = tmp_path / "accent.pptx"
    prs.save(path)
    assert "L2 强调色" in codes(deck_critique.critique(path))


def test_centered_title_with_left_body_is_review(tmp_path):
    prs = new_deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(slide, "居中标题", 30, center=True)
    add_text(slide, "左对齐正文", 15, y=1.9)
    add_filler(slide)
    path = tmp_path / "center.pptx"
    prs.save(path)
    report = deck_critique.critique(path)
    assert not report.failed
    assert any(item.rule == "L5 轴线" and item.level == "REVIEW" for item in report.findings)


def test_md_report_written(tmp_path):
    prs = new_deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(slide, "标题", 30)
    add_filler(slide)
    deck = tmp_path / "deck.pptx"
    prs.save(deck)
    md = tmp_path / "critique.md"
    code = deck_critique.main([str(deck), "--md", str(md)])
    assert code == 0
    text = md.read_text(encoding="utf-8")
    assert "页片审美评估" in text
    assert "人工三项" in text
