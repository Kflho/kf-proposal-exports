#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""deck_template.py — 按 PPT 模板生成：抽令牌（第 1 层）+ 套版式/复制页填内容（第 2 层）。

三件事：
  --tokens  读模板的主题色、中西文字体、画布、版式占位符、实际用到的字号/字体 → tokens.json
  --list   只看模板结构（有哪些版式、每页的占位符与形状），用来写 plan
  --plan   按 plan 生成：`layout` 页走占位符，`copy` 页复制模板某页再替换文字/图片

用法:
  python deck_template.py 模板.pptx --tokens out/tokens.json
  python deck_template.py 模板.pptx --list
  python deck_template.py 模板.pptx --plan plan.json --out out/deck.pptx

plan.json 片段:
  {"pages": [
     {"layout": "Title Slide", "title": "…", "subtitle": "…", "notes": "口播"},
     {"layout": 1, "title": "…", "body": ["要点一", "要点二"], "notes": "口播"},
     {"copy": 3, "text": {"原标题": "新标题"}, "images": {0: "figs/a.png"}}
  ]}

退出码: 0 成功；1 有未解析项（占位符/替换目标没找到）；2 参数或依赖错误。
"""
from __future__ import annotations

import argparse
import copy
import json
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

EMU = 914400.0
THEME = "ppt/theme/theme1.xml"
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
PLACEHOLDER_TYPES = {
    1: "title", 2: "body", 3: "center-title", 4: "subtitle", 5: "vertical-title",
    6: "vertical-body", 7: "object", 8: "chart", 9: "bitmap", 10: "media",
    11: "org-chart", 12: "table", 13: "slide-number", 14: "header",
    15: "footer", 16: "date", 17: "vertical-object", 18: "picture", 101: "slide-image",
}
KIND_ALIASES = {
    "title": ("title", "center-title"),
    "subtitle": ("subtitle",),
    "body": ("body", "object", "text", "vertical-body", "vertical-object"),
}
CHROME = {"date", "footer", "header", "slide-number", "slide-image"}


def _theme_tokens(prs) -> dict:
    from lxml import etree

    colors, fonts = {}, {}
    try:
        blob = prs.part.package.blob if hasattr(prs.part.package, "blob") else None
    except Exception:
        blob = None
    _ = blob
    # python-pptx 没有 theme API：从包关系里找 theme 部件
    for part in prs.part.package.iter_parts():
        if str(part.partname).endswith(THEME):
            root = etree.fromstring(part.blob)
            for node in root.findall(".//a:clrScheme/*", NS):
                name = etree.QName(node).localname
                srgb = node.find(".//a:srgbClr", NS)
                syscolor = node.find(".//a:sysClr", NS)
                value = srgb.get("val") if srgb is not None else (
                    syscolor.get("lastClr") if syscolor is not None else None)
                if value:
                    colors[name] = value.upper()
            for which in ("majorFont", "minorFont"):
                node = root.find(f".//a:{which}", NS)
                if node is None:
                    continue
                latin = node.find("a:latin", NS)
                ea = node.find("a:ea", NS)
                entry = {}
                if latin is not None and latin.get("typeface"):
                    entry["latin"] = latin.get("typeface")
                if ea is not None and ea.get("typeface"):
                    entry["ea"] = ea.get("typeface")
                if entry:
                    fonts[which] = entry
            break
    return {"colors": colors, "fonts": fonts}


def _placeholder_info(shape) -> dict:
    info = {"idx": None, "type": None, "name": shape.name}
    try:
        fmt = shape.placeholder_format
        if fmt is not None:
            info["idx"] = int(fmt.idx)
            info["type"] = PLACEHOLDER_TYPES.get(int(fmt.type), str(fmt.type))
    except Exception:
        pass
    return info


def _box(shape) -> dict:
    return {"x": round(shape.left / EMU, 2), "y": round(shape.top / EMU, 2),
            "w": round(shape.width / EMU, 2), "h": round(shape.height / EMU, 2)}


def _slide_stats(slide) -> dict:
    sizes, fonts, colors = set(), set(), set()
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                if run.font.size:
                    sizes.add(round(run.font.size.pt, 1))
                if run.font.name:
                    fonts.add(run.font.name)
                try:
                    rgb = run.font.color.rgb
                    if rgb is not None:
                        colors.add(str(rgb).upper())
                except Exception:
                    pass
    return {"sizes": sorted(sizes), "fonts": sorted(fonts), "colors": sorted(colors)}


def extract(prs) -> dict:
    layouts = []
    for index, layout in enumerate(prs.slide_layouts):
        placeholders = [_placeholder_info(ph) | _box(ph) for ph in layout.placeholders]
        layouts.append({"index": index, "name": layout.name, "placeholders": placeholders})
    slides = []
    for index, slide in enumerate(prs.slides, start=1):
        shapes = []
        for shape in slide.shapes:
            entry = {"kind": "picture" if shape.shape_type == 13 else
                     ("table" if getattr(shape, "has_table", False) else "shape"),
                     "name": shape.name, **_box(shape)}
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                entry["text"] = shape.text_frame.text[:80]
            if shape.shape_type == 13:
                entry["asset"] = Path(str(shape.image.filename or "image")).name
            shapes.append(entry)
        slides.append({"page": index, "layout": slide.slide_layout.name, "shapes": shapes,
                       **_slide_stats(slide)})
    sizes = sorted({size for slide in slides for size in slide["sizes"]})
    return {
        "slide_size": [round(prs.slide_width / EMU, 3), round(prs.slide_height / EMU, 3)],
        "theme": _theme_tokens(prs),
        "layouts": layouts,
        "slides": slides,
        "observed": {
            "sizes": sizes,
            "fonts": sorted({f for slide in slides for f in slide["fonts"]}),
            "colors": sorted({c for slide in slides for c in slide["colors"]}),
        },
    }


def _set_text(shape, value) -> None:
    frame = shape.text_frame
    if isinstance(value, list):
        first = True
        for item in value:
            para = frame.paragraphs[0] if first else frame.add_paragraph()
            first = False
            para.text = str(item)
        for para in (frame.paragraphs[len(value):] if not first else []):
            para._element.getparent().remove(para._element)
    else:
        frame.text = str(value)


def _pick(placeholders, kind: str | None, idx: int | None):
    if idx is not None:
        for ph in placeholders:
            if _placeholder_info(ph)["idx"] == idx:
                return ph
    if kind:
        accepted = KIND_ALIASES.get(kind, (kind,))
        for ph in placeholders:
            if _placeholder_info(ph)["type"] in accepted:
                return ph
    return None


def _replace_picture(shape, image_path: Path, slide) -> bool:
    try:
        image_part, r_id = slide.part.get_or_add_image(str(image_path))
        blip = shape._element.blipFill.blip
        blip.embed = r_id
        return True
    except Exception:
        return False


def _clear_slides(prs) -> int:
    """模板里的样例页先清掉，只留版式；plan 决定最终页序。"""
    from pptx.oxml.ns import qn

    lst = prs.slides._sldIdLst
    removed = 0
    for sld in list(lst):
        r_id = sld.get(qn("r:id"))
        lst.remove(sld)
        try:
            prs.part.drop_rel(r_id)
        except Exception:
            pass
        removed += 1
    return removed


def build(template_path: Path, plan: dict, out_path: Path) -> dict:
    from pptx import Presentation

    prs = Presentation(str(template_path))
    # 复制页要先在清空之前把源页的形状深拷出来
    clones: dict[int, tuple] = {}
    for page_no, page in enumerate(plan.get("pages", []), start=1):
        if "copy" in page:
            source = prs.slides[int(page["copy"]) - 1]
            clones[page_no] = (source.slide_layout,
                               [copy.deepcopy(sh._element) for sh in source.shapes])
    _clear_slides(prs)
    unresolved = []
    for page_no, page in enumerate(plan.get("pages", []), start=1):
        notes = page.get("notes", "")
        if "copy" in page:
            layout, elements = clones[page_no]
            slide = prs.slides.add_slide(layout)
            for shape in list(slide.shapes):
                shape._element.getparent().remove(shape._element)
            for element in elements:
                slide.shapes._spTree.append(copy.deepcopy(element))
            text_map = page.get("text") or {}
            for old, new in text_map.items():
                hit = False
                for shape in slide.shapes:
                    if not getattr(shape, "has_text_frame", False):
                        continue
                    for para in shape.text_frame.paragraphs:
                        for run in para.runs:
                            if old in run.text:
                                run.text = run.text.replace(old, new)
                                hit = True
                if not hit:
                    unresolved.append({"page": page_no, "kind": "text", "target": old})
            for index_key, image_path in (page.get("images") or {}).items():
                index = int(index_key)
                pictures = [s for s in slide.shapes if s.shape_type == 13]
                if index >= len(pictures):
                    unresolved.append({"page": page_no, "kind": "image", "target": index})
                    continue
                if not _replace_picture(pictures[index], Path(image_path), slide):
                    unresolved.append({"page": page_no, "kind": "image", "target": index})
        else:
            layout_ref = page.get("layout", 0)
            if isinstance(layout_ref, int):
                layout = prs.slide_layouts[layout_ref]
            else:
                layout = next((l for l in prs.slide_layouts if l.name == layout_ref), None)
            if layout is None:
                unresolved.append({"page": page_no, "kind": "layout", "target": layout_ref})
                continue
            slide = prs.slides.add_slide(layout)
            for kind, value in (("title", page.get("title")),
                                ("subtitle", page.get("subtitle")),
                                ("body", page.get("body"))):
                if value is None:
                    continue
                if kind == "title" and slide.shapes.title is not None:
                    shape = slide.shapes.title
                else:
                    shape = _pick(slide.placeholders, kind, page.get(f"{kind}_idx"))
                if shape is None:
                    unresolved.append({"page": page_no, "kind": "placeholder", "target": kind})
                    continue
                _set_text(shape, value)
            for key, value in (page.get("placeholders") or {}).items():
                shape = _pick(slide.placeholders, None, int(key))
                if shape is None:
                    unresolved.append({"page": page_no, "kind": "placeholder", "target": key})
                    continue
                _set_text(shape, value)
            for index_key, image_path in (page.get("images") or {}).items():
                shape = _pick(slide.placeholders, None, int(index_key))
                if shape is None:
                    unresolved.append({"page": page_no, "kind": "image", "target": index_key})
                    continue
                try:
                    shape.insert_picture(str(image_path))
                except Exception:
                    unresolved.append({"page": page_no, "kind": "image", "target": index_key})
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return {"out": str(out_path), "pages": len(plan.get("pages", [])),
            "unresolved": unresolved}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="按 PPT 模板出片：抽令牌 / 看结构 / 套版式生成")
    parser.add_argument("template", help="模板 .pptx")
    parser.add_argument("--tokens", metavar="JSON", help="抽令牌写到该文件")
    parser.add_argument("--list", action="store_true", help="只打印结构（版式/页/形状）")
    parser.add_argument("--plan", metavar="JSON", help="内容 plan")
    parser.add_argument("--out", help="输出 .pptx（配合 --plan）")
    parser.add_argument("--json", action="store_true", help="输出 JSON")
    args = parser.parse_args(argv)

    template = Path(args.template)
    if not template.is_file():
        print(f"错误：找不到 {template}", file=sys.stderr)
        return 2
    try:
        from pptx import Presentation

        prs = Presentation(str(template))
        if args.list:
            tokens = extract(prs)
            if args.json:
                print(json.dumps(tokens, ensure_ascii=False, indent=1))
            else:
                print(f"{template}：{tokens['slide_size'][0]}x{tokens['slide_size'][1]} in，"
                      f"{len(tokens['layouts'])} 个版式，{len(tokens['slides'])} 页")
                for layout in tokens["layouts"]:
                    items = [p for p in layout["placeholders"] if p["type"] not in CHROME]
                    phs = "，".join(f"{p['type']}#{p['idx']}" for p in items) or "无占位符"
                    print(f"  [{layout['index']}] {layout['name']}：{phs}")
                print("  观察到的字号：", tokens["observed"]["sizes"])
                print("  观察到的字体：", tokens["observed"]["fonts"])
            return 0
        if args.tokens:
            tokens = extract(prs)
            Path(args.tokens).parent.mkdir(parents=True, exist_ok=True)
            Path(args.tokens).write_text(json.dumps(tokens, ensure_ascii=False, indent=1),
                                         encoding="utf-8")
            print(f"{template} → {args.tokens}（{len(tokens['layouts'])} 版式 / "
                  f"{len(tokens['slides'])} 页 / 主题色 {len(tokens['theme']['colors'])} 个）")
            return 0
        if args.plan:
            if not args.out:
                print("错误：--plan 需要配合 --out", file=sys.stderr)
                return 2
            plan = json.loads(Path(args.plan).read_text(encoding="utf-8-sig"))
            result = build(template, plan, Path(args.out))
            if args.json:
                print(json.dumps(result, ensure_ascii=False, indent=1))
            else:
                print(f"生成 {result['out']}：{result['pages']} 页，"
                      f"未解析 {len(result['unresolved'])} 项")
                for item in result["unresolved"]:
                    print(f"  ! P{item['page']} {item['kind']}：{item['target']}")
            return 0 if not result["unresolved"] else 1
        parser.error("需要 --tokens / --list / --plan 之一")
    except ModuleNotFoundError as exc:
        print(f"错误：缺少依赖（{exc.name}）：pip install python-pptx Pillow", file=sys.stderr)
        return 2
    except Exception as exc:  # noqa: BLE001 - CLI 边界
        print(f"错误：{exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
