#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""text_check.py — 页片/报告的文字自检。

规则数据在 ``assets/text/banned.json``：
  - ``words`` / ``allow_context``：禁用词与放行语境；
  - ``patterns``：套话开头；
  - ``ai_patterns``：翻译腔与空泛句式；
  - ``verdict_markers``：标题里的判断词。

判据见 references/text-quality.md。退出码 0 无 error / 1 有 error / 2 环境错误。
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import asdict, dataclass, field
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

ROOT = Path(__file__).resolve().parents[1]
RULES_PATH = ROOT / "assets" / "text" / "banned.json"
TOPIC_WORDS = ("概述", "介绍", "分析", "方法", "模型", "结果", "总结", "讨论", "方案", "背景", "结论")
SENTENCE_SPLIT = re.compile(r"[。！？!?；;\n]+")
PAGE_NUM_RE = re.compile(r"^\d+\s*/\s*\d+$")


@dataclass
class Issue:
    level: str
    code: str
    page: int
    message: str


@dataclass
class Report:
    path: str
    pages: int = 0
    errors: list[Issue] = field(default_factory=list)
    warnings: list[Issue] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors


def load_rules(path: str | Path | None = None) -> dict:
    """读取词表；路径默认是 skill 的 assets/text/banned.json。"""
    target = Path(path) if path else RULES_PATH
    data = json.loads(target.read_text(encoding="utf-8-sig"))
    return {
        "words": list(data.get("words") or []),
        "allow_context": dict(data.get("allow_context") or {}),
        "patterns": [re.compile(p) for p in (data.get("patterns") or [])],
        "ai_patterns": [re.compile(p) for p in (data.get("ai_patterns") or [])],
        "verdict_markers": list(data.get("verdict_markers") or []),
    }


def _banned_hit(text: str, rules: dict) -> str | None:
    for word in rules["words"]:
        if word not in text:
            continue
        allowed = rules["allow_context"].get(word) or []
        if any(context in text for context in allowed):
            continue
        return word
    return None


def _pattern_hit(text: str, patterns) -> tuple[str, str] | None:
    for pattern in patterns:
        match = pattern.search(text)
        if match:
            return pattern.pattern, match.group(0)
    return None


def _cjk_len(text: str) -> int:
    """按"一个汉字算 1、一个西文单词算 1"估长度，避免英文标题被误判过长。"""
    units = 0
    latin_run = 0
    for ch in text:
        if ch.isascii() and (ch.isalnum() or ch in "+-/"):
            latin_run += 1
        else:
            if latin_run:
                units += 1
                latin_run = 0
            if ch.strip():
                units += 1
    if latin_run:
        units += 1
    return units


def _norm(text: str) -> str:
    return re.sub(r"[\s，。：:、；;！？!?·—\-（）()【】\[\]]+", "", text)


def _dedupe(issues: list[Issue]) -> list[Issue]:
    seen, out = set(), []
    for issue in issues:
        key = (issue.code, issue.page, issue.message)
        if key in seen:
            continue
        seen.add(key)
        out.append(issue)
    return out


def _load_pptx(path: Path):
    from pptx import Presentation

    prs = Presentation(str(path))
    pages = []
    for index, slide in enumerate(prs.slides, start=1):
        frames = []
        for shape in slide.shapes:
            if not shape.has_text_frame or shape.has_table:
                continue
            text = shape.text_frame.text.strip()
            if not text or PAGE_NUM_RE.match(text):
                continue
            size = 0
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.size and run.font.size.pt > size:
                        size = run.font.size.pt
            frames.append((size, shape.text_frame))
        if frames:
            def is_title_like(text: str) -> bool:
                t = text.strip()
                return len(t) >= 6 and not re.fullmatch(r"[\d.%+\-/\s]+", t)

            preferred = [f for f in frames if is_title_like(f[1].text)]
            pool = preferred or frames
            pool.sort(key=lambda item: item[0], reverse=True)
            title_frame = pool[0][1]
            title = title_frame.text.strip().split("\n")[0]
            body = frames[1:]
        else:
            title, body = "", []
        bullets, sentences = [], []
        for size, frame in body:
            if size and not (10 <= size <= 22):
                continue
            for para in frame.paragraphs:
                line = "".join(r.text for r in para.runs).strip() or para.text.strip()
                if line and len(line) >= 4 and not re.fullmatch(r"[\d.%+\-/\s]+", line):
                    bullets.append(line)
        for _, frame in frames:
            sentences.extend(s for s in SENTENCE_SPLIT.split(frame.text) if s.strip())
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        pages.append({"page": index, "title": title, "bullets": bullets,
                      "sentences": sentences, "notes": notes})
    return pages


def _load_markdown(path: Path):
    pages = []
    current = {"page": 1, "title": "", "bullets": [], "sentences": [], "notes": ""}
    in_code = False
    for raw in path.read_text(encoding="utf-8-sig").splitlines():
        line = raw.rstrip()
        if line.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            continue
        if line.startswith("#"):
            if current["title"] or current["bullets"] or current["sentences"]:
                pages.append(current)
                current = {"page": len(pages) + 1, "title": "", "bullets": [],
                           "sentences": [], "notes": ""}
            current["title"] = line.lstrip("#").strip()
            continue
        if re.match(r"^\s*[-*+]\s+", line) or re.match(r"^\s*\d+[.)]\s+", line):
            current["bullets"].append(re.sub(r"^\s*(?:[-*+]|\d+[.)])\s+", "", line).strip())
            continue
        if line.strip():
            current["sentences"].extend(s for s in SENTENCE_SPLIT.split(line) if s.strip())
    if current["title"] or current["bullets"] or current["sentences"]:
        pages.append(current)
    return pages


def check_pages(pages, rules=None, max_title=30, max_sentence=60, max_bullets=5,
                max_bullet=40, max_notes=120) -> Report:
    rules = rules or load_rules()
    report = Report(path="", pages=len(pages))
    for item in pages:
        page = item["page"]

        def add(level, code, message):
            (report.errors if level == "error" else report.warnings).append(
                Issue(level=level, code=code, page=page, message=message))

        title = item.get("title", "").strip()
        if title:
            hit = _banned_hit(title, rules)
            if hit:
                add("error", "banned-phrase", f"标题禁用词「{hit}」：{title[:30]}")
            phrase = _pattern_hit(title, rules["ai_patterns"])
            if phrase:
                add("warning", "ai-phrase", f"标题句式「{phrase[1][:20]}」：{title[:30]}")

            norm = _norm(title)
            topic = False
            if _cjk_len(title) > max_title:
                add("warning", "title-too-long",
                    f"标题 {_cjk_len(title)} 字（≤{max_title}）：{title[:30]}")
            elif _cjk_len(norm) <= 6 or (norm.endswith(TOPIC_WORDS) and len(norm) <= 12):
                add("warning", "title-topic", f"标题像话题不像结论：{title[:30]}")
                topic = True
            elif any(w in norm for w in TOPIC_WORDS) and not re.search(r"[，,：:；;]", title):
                if _cjk_len(norm) <= 12:
                    add("warning", "title-topic", f"标题像话题不像结论：{title[:30]}")
                    topic = True
            if not topic and not any(marker in title for marker in rules["verdict_markers"]):
                add("warning", "title-no-verdict",
                    f"标题没有判断或数字，像话题：{title[:30]}")

        def check_text(text: str, where: str) -> None:
            hit = _banned_hit(text, rules)
            if hit:
                add("error", "banned-phrase", f"{where}禁用词「{hit}」：{text[:24]}")
            phrase = _pattern_hit(text, rules["ai_patterns"])
            if phrase:
                add("warning", "ai-phrase", f"{where}句式「{phrase[1][:20]}」：{text[:24]}")
            pattern = _pattern_hit(text, rules["patterns"])
            if pattern:
                add("error", "banned-phrase", f"{where}套话开头「{pattern[1][:20]}」：{text[:24]}")

        for sentence in item.get("sentences", []):
            check_text(sentence, "句子")
            if _cjk_len(sentence) > max_sentence:
                add("warning", "sentence-too-long",
                    f"单句 {_cjk_len(sentence)} 字（≤{max_sentence}）：{sentence[:24]}")

        bullets = item.get("bullets", [])
        if len(bullets) > max_bullets:
            add("warning", "too-many-bullets", f"要点 {len(bullets)} 条（≤{max_bullets}）")
        for bullet in bullets:
            check_text(bullet, "要点")
            if _cjk_len(bullet) > max_bullet:
                add("warning", "bullet-too-long",
                    f"要点 {_cjk_len(bullet)} 字（≤{max_bullet}）：{bullet[:24]}")

        notes = item.get("notes", "")
        if notes:
            check_text(notes, "口播")
        if _cjk_len(notes) > max_notes:
            add("warning", "notes-too-long", f"口播 {_cjk_len(notes)} 字（≤{max_notes}）")

    report.errors = _dedupe(report.errors)
    report.warnings = _dedupe(report.warnings)
    return report


def check_file(path, rules=None, **kwargs) -> Report:
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(path)
    if path.suffix.lower() in {".pptx", ".potx"}:
        pages = _load_pptx(path)
    elif path.suffix.lower() in {".md", ".markdown", ".txt"}:
        pages = _load_markdown(path)
    else:
        raise ValueError(f"不支持的格式：{path.suffix}（.pptx / .md）")
    report = check_pages(pages, rules=rules, **kwargs)
    report.path = str(path)
    return report


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="页片/报告文字自检")
    parser.add_argument("path", help=".pptx 或 .md")
    parser.add_argument("--json", action="store_true")
    parser.add_argument("--rules", default=str(RULES_PATH), help="词表 JSON")
    parser.add_argument("--max-title", type=int, default=30)
    parser.add_argument("--max-sentence", type=int, default=60)
    parser.add_argument("--max-bullets", type=int, default=5)
    parser.add_argument("--max-bullet", type=int, default=40)
    parser.add_argument("--max-notes", type=int, default=120)
    args = parser.parse_args(argv)

    try:
        rules = load_rules(args.rules)
        report = check_file(args.path, rules=rules, max_title=args.max_title,
                            max_sentence=args.max_sentence, max_bullets=args.max_bullets,
                            max_bullet=args.max_bullet, max_notes=args.max_notes)
    except ModuleNotFoundError as exc:
        print(f"错误：缺少依赖（{exc.name}）：pip install python-pptx", file=sys.stderr)
        return 2
    except Exception as exc:  # noqa: BLE001 - CLI 边界
        print(f"错误：{exc}", file=sys.stderr)
        return 2

    if args.json:
        payload = asdict(report)
        payload["ok"] = report.ok
        print(json.dumps(payload, indent=1, ensure_ascii=False))
    else:
        print(f"{report.path}：{report.pages} 页/节")
        for issue in report.errors + report.warnings:
            mark = "E" if issue.level == "error" else "W"
            print(f"  [{mark}] §{issue.page} [{issue.code}] {issue.message}")
        print(f"summary: {len(report.errors)} error / {len(report.warnings)} warning")
    return 0 if report.ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
